# -*- coding: utf-8 -*-
"""Génère presentation_finale.pptx (14 slides + annexes, notes orateur)."""
import re
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).parent          # le script vit dans presentation/
A = ROOT / "assets"

# --- palette (cohérente avec le poster) ---
NAVY   = RGBColor(0x16, 0x32, 0x4F)
GOLD   = RGBColor(0xFF, 0xD1, 0x66)
RED    = RGBColor(0xC0, 0x39, 0x2B)
GREEN  = RGBColor(0x1E, 0x7D, 0x4F)
GREY   = RGBColor(0x55, 0x5F, 0x6B)
LIGHT  = RGBColor(0xEE, 0xF3, 0xF8)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DARK   = RGBColor(0x1A, 0x1A, 0x2E)

SW, SH = Inches(13.333), Inches(7.5)
prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]

def img_path(prefix):
    hits = sorted(A.glob(prefix + "*"))
    assert hits, f"figure absente: {prefix}"
    return str(hits[0])

def add_rich(p, text, size, color=DARK, bold_color=None):
    """Écrit `text` dans le paragraphe p, **gras** parsé."""
    p.text = ""
    for i, seg in enumerate(re.split(r"\*\*", text)):
        if not seg:
            continue
        r = p.add_run(); r.text = seg
        r.font.size = Pt(size); r.font.name = "Calibri"
        if i % 2 == 1:
            r.font.bold = True
            r.font.color.rgb = bold_color or color
        else:
            r.font.color.rgb = color
    return p

def new_slide(title, subtitle=None, notes=None):
    s = prs.slides.add_slide(BLANK)
    # bandeau titre
    bar = s.shapes.add_shape(1, 0, 0, SW, Inches(1.0))  # 1 = rectangle
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    tf = bar.text_frame; tf.margin_left = Inches(0.45); tf.margin_right = Inches(0.3)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = True
    add_rich(tf.paragraphs[0], title, 26, color=WHITE, bold_color=GOLD)
    if subtitle:
        p2 = tf.add_paragraph(); add_rich(p2, subtitle, 12, color=RGBColor(0xC9, 0xD6, 0xE3))
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s

def bullets(s, items, left, top, width, height, size=15, gap=6):
    tb = s.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame; tf.word_wrap = True
    first = True
    for it in items:
        lvl = 0
        if isinstance(it, tuple):
            it, lvl = it
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap); p.level = lvl
        add_rich(p, ("• " if lvl == 0 else "– ") + it, size - 2 * lvl, bold_color=NAVY)
    return tb

def picture(s, path, left, top, width=None, height=None):
    if width and not height:
        w, h = Image.open(path).size
        height = Emu(int(width * h / w))
    return s.shapes.add_picture(path, left, top, width=width, height=height)

def caption(s, text, left, top, width, size=10):
    tb = s.shapes.add_textbox(left, top, width, Inches(0.35))
    p = tb.text_frame.paragraphs[0]
    add_rich(p, text, size, color=GREY)
    p.runs and setattr(p.runs[0].font, "italic", True)
    return tb

def chip(s, big, small, left, top, w=Inches(2.6), h=Inches(1.25), color=NAVY):
    box = s.shapes.add_shape(5, left, top, w, h)  # 5 = rounded rect
    box.fill.solid(); box.fill.fore_color.rgb = color; box.line.fill.background()
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = big; r.font.size = Pt(26); r.font.bold = True
    r.font.color.rgb = GOLD; r.font.name = "Calibri"
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = small; r2.font.size = Pt(10.5); r2.font.color.rgb = WHITE
    return box

def table(s, rows, left, top, width, col_widths=None, font=12, header_fill=NAVY):
    n_r, n_c = len(rows), len(rows[0])
    shp = s.shapes.add_table(n_r, n_c, left, top, width, Inches(0.4 * n_r)).table
    if col_widths:
        total = sum(col_widths)
        for j, cw in enumerate(col_widths):
            shp.columns[j].width = Emu(int(width * cw / total))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = shp.cell(i, j); c.text = ""
            p = c.text_frame.paragraphs[0]
            add_rich(p, str(val), font if i else font, bold_color=NAVY,
                     color=WHITE if i == 0 else DARK)
            if i == 0:
                p.runs and setattr(p.runs[0].font, "bold", True)
                c.fill.solid(); c.fill.fore_color.rgb = header_fill
            else:
                c.fill.solid()
                c.fill.fore_color.rgb = WHITE if i % 2 else RGBColor(0xF3, 0xF6, 0xFA)
            c.margin_top = c.margin_bottom = Pt(3)
    return shp

def band(s, text, left, top, width, height=Inches(0.85), fill=LIGHT, accent=NAVY, size=13):
    box = s.shapes.add_shape(1, left, top, width, height)
    box.fill.solid(); box.fill.fore_color.rgb = fill; box.line.color.rgb = accent
    box.line.width = Pt(1)
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.18); tf.margin_right = Inches(0.15)
    add_rich(tf.paragraphs[0], text, size, bold_color=accent)
    return box

# S1 titre
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(1, 0, 0, SW, SH)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
tb = s.shapes.add_textbox(Inches(0.7), Inches(1.5), Inches(7.2), Inches(3.2))
tf = tb.text_frame; tf.word_wrap = True
add_rich(tf.paragraphs[0], "Réduire les accidents par la **rénovation d'infrastructure**", 36, color=WHITE, bold_color=GOLD)
p = tf.add_paragraph(); add_rich(p, "\nPipeline d'aide à la décision pour une agence de voirie (DOT)", 20, color=WHITE)
p = tf.add_paragraph(); add_rich(p, "Données US Accidents - 7,73 M accidents , 2016-2023 , 49 États", 14, color=RGBColor(0xC9, 0xD6, 0xE3))
p = tf.add_paragraph(); add_rich(p, "\nSAÉ DataValueXploring - juillet 2026", 13, color=RGBColor(0xC9, 0xD6, 0xE3))
picture(s, img_path("00_carte"), Inches(8.15), Inches(1.15), width=Inches(4.6))
caption(s, "50 zones prioritaires - Californie (démo)", Inches(8.15), Inches(4.55), Inches(4.6), 10)
for i, (b, sm) in enumerate([("7,73 M", "accidents analysés"), ("×62", "concentration du risque"),
                              ("ρ = 0,78", "stabilité temporelle"), ("6 860", "graves évitables (est.)")]):
    chip(s, b, sm, Inches(0.7 + i * 3.05), Inches(5.6), w=Inches(2.75), h=Inches(1.15),
         color=RGBColor(0x1F, 0x42, 0x63))
s.notes_slide.notes_text_frame.text = ("Accroche : un gestionnaire de voirie a un budget limité et des milliers de "
    "carrefours. On lui dit lesquels rénover, quoi y installer, et ce que ça devrait rapporter.")

# S2 problème
s = new_slide("Le problème métier", "Rôle incarné : agence de voirie / DOT d'une collectivité",
    notes="Insister : le choix du rôle rend le projet actionnable, vs analyses génériques (« la pluie est dangereuse ») qui n'aident aucun décideur.")
tb = s.shapes.add_textbox(Inches(0.7), Inches(1.4), Inches(11.9), Inches(1.1))
add_rich(tb.text_frame.paragraphs[0],
         "« Comment réduire durablement les accidents sur notre réseau, avec un **budget de travaux limité** ? »",
         22, bold_color=RED)
bullets(s, [
    "Aujourd'hui : priorisation **réactive** - remontées terrain, accidents médiatisés, arbitrages politiques",
    "Notre apport : une priorisation **data-driven, explicable et défendable**",
    "Périmètre d'action du rôle : **l'infrastructure** (feux, passages piétons, stops, modération de trafic)",
    ("on n'agit ni sur la météo, ni sur le comportement des conducteurs", 1),
], Inches(0.8), Inches(2.7), Inches(11.7), Inches(3.2), size=17, gap=12)
band(s, "La donnée crée de la valeur quand elle répond à une décision : **où investir le prochain euro de travaux ?**",
     Inches(0.7), Inches(6.2), Inches(11.9), accent=NAVY, size=15)

# S3 données
s = new_slide("Les données - et leurs pièges", "US Accidents (Kaggle) , Feb 2016 -> Mar 2023",
    notes="Montrer qu'on connaît les limites AVANT de montrer des résultats : crédibilité. Biais de reporting = collecte via API trafic MapQuest/Bing.")
bullets(s, [
    "**7 728 394 accidents**, 46 colonnes, 49 États",
    "Collecte via API trafic -> **biais de reporting assumé** : peu de lignes ≠ peu d'accidents",
    "Top 3 États ≈ **41 %** du volume (CA, FL, TX)",
    "Conséquence méthodologique : on ne compare **jamais** les États entre eux -",
    ("on raisonne à l'intérieur d'un territoire (analyse paramétrée par une variable STATE)", 1),
], Inches(0.7), Inches(1.45), Inches(5.6), Inches(4.6), size=15, gap=10)
picture(s, img_path("02_a1"), Inches(6.5), Inches(1.6), width=Inches(6.2))
caption(s, "Couverture très inégale : top 15 États par volume", Inches(6.5), Inches(4.35), Inches(6.2))

# S4 cadrage
s = new_slide("Le cadrage : 3 questions, 3 natures", "Pas de modèle « par habitude » - la méthode suit la nature de la question",
    notes="LA slide méthodo. Question attendue : « pourquoi un prédicteur si vous ne prédisez pas ? » -> 1) tâche évaluable, 2) contrôle des confondants, 3) la convergence entre méthodes devient une preuve.")
table(s, [
    ["Question métier", "Nature", "Méthode"],
    ["OÙ rénover en priorité ?", "concentration -> descriptif", "agrégation H3 + exposition OSM"],
    ["QUOI installer sur place ?", "attribution -> ML (moteur de preuve)", "modèle de gravité -> effets ajustés"],
    ["QUAND agir temporairement ?", "conjoncturel -> descriptif", "patterns temporels / météo"],
], Inches(0.7), Inches(1.5), Inches(11.9), col_widths=[3.4, 3.6, 4.9], font=14)
band(s, "**La prédiction est le moyen ; l'explication est le livrable.** On entraîne un modèle de gravité pour "
        "l'interroger - quels équipements aggravent, à contexte égal - pas pour prédire.",
     Inches(0.7), Inches(3.75), Inches(11.9), height=Inches(1.0), accent=RED, size=15)
bullets(s, [
    "**OÙ** est une question de concentration -> agrégation, pas de ML supervisé (assumé)",
    "**QUOI** est une question d'attribution -> c'est là que le ML apporte de la valeur",
    "**QUAND** est conjoncturel -> visualisé, pas modélisé (pour ne pas diluer le fil rouge)",
], Inches(0.8), Inches(5.05), Inches(11.7), Inches(1.9), size=14, gap=8)

# S5 architecture
s = new_slide("Architecture de la pipeline", "Un notebook unique, exécutable de bout en bout sans intervention",
    notes="Mentionner les tests : assertions bloquantes dans le notebook (intégrité + anti-fuite) + suite pytest sur le livrable + Docker.")
mono = ("Setup reproductible  ->  Chargement Spark (schéma explicite)  ->  Nettoyage explicite & testé\n"
        "   ├─  A. EDA : volumétrie, biais, sévérité, « QUAND »\n"
        "   ├─  B. OÙ  : grille H3 -> charge pondérée gravité -> exposition OSM (acc/km) + stabilité\n"
        "   ├─  C. QUOI: 5 modèles comparés -> effets ajustés (odds ratios + g-computation) -> leviers\n"
        "   ├─  D. Synthèse décisionnelle : par zone, où + quoi + impact estimé  (CSV + carte)\n"
        "   └─  E. Analyse critique : ce qui marche / pas, limites, décision terrain")
box = s.shapes.add_shape(1, Inches(0.7), Inches(1.5), Inches(11.9), Inches(2.5))
box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xF6, 0xF8, 0xFA); box.line.color.rgb = NAVY
tf = box.text_frame; tf.word_wrap = True; tf.margin_left = Inches(0.25)
for i, line in enumerate(mono.split("\n")):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    r = p.add_run(); r.text = line; r.font.name = "Consolas"; r.font.size = Pt(13); r.font.color.rgb = DARK
bullets(s, [
    "**PySpark** pour les 7,73 M lignes (chargement, nettoyage, agrégations, split) - pandas saturerait",
    "**scikit-learn sur échantillon maîtrisé** (800 k, plafond documenté) pour la comparaison de modèles",
    "**Reproductible** : données auto-téléchargées, seeds fixés, versions figées, tests, Docker",
], Inches(0.8), Inches(4.3), Inches(11.7), Inches(2.3), size=15, gap=10)

# S6 OÙ concentration
s = new_slide("OÙ - le risque est très concentré",
    "Californie (démonstration) : 1,74 M accidents -> 58 857 zones H3 de ~0,7 km²",
    notes="Charge = somme des sévérités = volume pondéré gravité. 100 accidents légers ≠ 100 graves. H3 : cellules équi-surface, rés. 8 = échelle d'un carrefour.")
picture(s, img_path("08_b_"), Inches(0.55), Inches(1.7), width=Inches(8.3))
tb = bullets(s, [
    "**50 zones** (0,08 %) portent **5,3 %** de la charge totale",
    "soit **×62** leur « part équitable »",
    "Rénover quelques carrefours a un effet de levier disproportionné",
], Inches(9.1), Inches(2.0), Inches(3.7), Inches(3.4), size=15, gap=12)
band(s, "On ne saupoudre pas le budget : on concentre les travaux là où la charge s'accumule.",
     Inches(9.1), Inches(5.3), Inches(3.7), height=Inches(1.3), accent=RED, size=13)

# S7 OÙ validations
s = new_slide("OÙ - deux validations qui rendent le classement défendable",
    notes="Question attendue « pourquoi H3 ? » : cellules équi-surface (comptage = densité comparable), voisinage hexagonal régulier, rés. 8 ≈ carrefour.")
picture(s, img_path("09_b2"), Inches(0.55), Inches(1.55), width=Inches(5.2))
caption(s, "Score par zone : 2016-2019 vs 2020-2023", Inches(0.55), Inches(6.7), Inches(5.2))
bullets(s, [
    "**1. Stabilité temporelle** - les points noirs sont-ils du bruit ?",
    ("corrélation de rang **Spearman ρ = 0,776** entre les deux périodes", 1),
    ("les **50 zones prioritaires restent toutes dans le top 5 %** -> risque structurel", 1),
    "**2. Exposition routière (OSM)** - dangereux, ou juste fréquenté ?",
    ("longueur de voirie par zone -> **accidents/km** -> re-classement", 1),
    ("une zone gagne **+11 rangs** une fois normalisée : point noir concentré masqué par le volume", 1),
    ("couverture : 42/50 zones (API Overpass limitée - cache local livré, dégradation documentée)", 1),
], Inches(6.1), Inches(1.6), Inches(6.7), Inches(5.4), size=15, gap=9)

# S8 comparaison modèles
s = new_slide("QUOI - une comparaison honnête de 5 modèles",
    "Split temporel : train < 2022, test 2022-2023 - on prédit l'avenir avec le passé",
    notes="Performance modeste ASSUMÉE : la gravité dépend de facteurs absents (vitesse réelle, alcool). Notre critère de succès = fiabilité de l'explication, pas la perf brute - c'est mot pour mot dans le sujet. Pourquoi pas SMOTE : class_weight, on ne fabrique pas de données.")
table(s, [
    ["Modèle", "PR-AUC", "ROC-AUC", "Rappel « grave »"],
    ["Dummy (plancher)", "0,064", "0,500", "0,00"],
    ["Régression logistique", "0,110", "0,635", "0,63"],
    ["Arbre de décision", "0,079", "0,560", "0,68"],
    ["Random Forest", "0,115", "0,646", "0,63"],
    ["**HistGradientBoosting** ✓", "**0,125**", "**0,676**", "0,57"],
], Inches(0.7), Inches(1.55), Inches(7.6), col_widths=[3.4, 1.4, 1.4, 1.7], font=13)
bullets(s, [
    "**Pas d'accuracy** : à 80/20, prédire « non-grave » partout = 80 % sans rien apprendre",
    "PR-AUC : la métrique honnête quand la classe positive est rare (19,5 %)",
    "Tous battent le plancher ; le retenu fait **×2 vs plancher**",
    "Déséquilibre géré par **class_weight** (pas de données fabriquées)",
], Inches(8.6), Inches(1.7), Inches(4.3), Inches(4.4), size=13.5, gap=10)
band(s, "Le modèle a appris un **signal réel** - c'est tout ce qu'on lui demande : il sert de moteur de preuve, pas d'oracle.",
     Inches(0.7), Inches(6.15), Inches(11.9), accent=NAVY, size=14)

# S9 effets ajustés
s = new_slide("QUOI - les effets ajustés : la réponse claire",
    "Deux lectures indépendantes, une même question : effet net de chaque équipement, à contexte égal ?",
    notes="Point clé : une moyenne brute « avec/sans feux » dirait l'INVERSE (les feux sont là où c'est dense). Les effets ajustés renversent la lecture - c'est LA démonstration de l'utilité du modèle. SHAP en annexe : requête d'intervention (g-comp) ≠ requête d'attribution (SHAP).")
picture(s, img_path("11_c3"), Inches(0.45), Inches(1.6), width=Inches(6.3))
picture(s, img_path("12_c3"), Inches(6.95), Inches(1.6), width=Inches(6.0))
bullets(s, [
    "**Odds ratios (IC 95 %)** : passage piéton **0,41** , stop **0,27** , feux **0,39** -> protecteurs ; carrefour 1,2 , voie ferrée 1,7 -> aggravants",
    "**G-computation** (modèle retenu) : on « allume/éteint » l'équipement à contexte constant -> passage piéton **−20 pts**, stop **−19 pts** de probabilité de gravité",
    "**Concordance des deux lectures : 10/12 équipements** -> diagnostic robuste",
], Inches(0.7), Inches(5.05), Inches(11.9), Inches(2.2), size=14, gap=8)

# S10 prescription
s = new_slide("QUOI - de l'explication à la prescription",
    notes="Question attendue « c'est causal ? » : NON - associations ajustées. L'impact est ex-ante, il sert à PRIORISER. Mesure réelle = pilote + différence de différences (slide 13).")
mono = ("pour chaque zone prioritaire :\n"
        "   aléa local dominant (lift)              ex. « carrefour » sur-représenté ×2,3\n"
        "   × mesure protectrice MANQUANTE          ex. pas de stop dans la zone\n"
        "   × effet ajusté conditionnel du modèle   Δ_adj(stop | carrefour) = −19,9 pts\n"
        "   -> recommandation + impact estimé = |Δ_adj| × volume  =  accidents GRAVES évités")
box = s.shapes.add_shape(1, Inches(0.7), Inches(1.45), Inches(11.9), Inches(2.15))
box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xF6, 0xF8, 0xFA); box.line.color.rgb = NAVY
tf = box.text_frame; tf.word_wrap = True; tf.margin_left = Inches(0.25)
for i, line in enumerate(mono.split("\n")):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    r = p.add_run(); r.text = line; r.font.name = "Consolas"; r.font.size = Pt(13.5); r.font.color.rgb = DARK
bullets(s, [
    "Exemple réel : zone **#2** (2 543 accidents, aléa carrefour) -> « stop / cédez-le-passage », **≈ 505 graves évités** (ex-ante)",
    "**6 types** de recommandations sur les 50 zones - l'effet conditionnel évite le « même aménagement partout »",
    "13 zones sans levier infra dominant -> « **investiguer facteur humain** » (affiché honnêtement, pas masqué)",
    "Contre-vérification systématique : effet ajusté (modèle) <-> effet brut (données)",
], Inches(0.8), Inches(3.95), Inches(11.7), Inches(2.9), size=15, gap=10)

# S11 QUAND
s = new_slide("QUAND - les contre-intuitifs qui font gagner du budget",
    "Risque conjoncturel ≠ risque structurel",
    notes="Hypothèse brouillard : conduite plus prudente - on le dit comme hypothèse, pas comme fait. La valeur différenciante = contredire le common knowledge avec des données.")
picture(s, img_path("06_a3"), Inches(0.55), Inches(1.6), width=Inches(5.9))
picture(s, img_path("07_a4"), Inches(6.75), Inches(1.6), width=Inches(6.1))
bullets(s, [
    "**Le brouillard** affiche la part d'accidents graves **la plus basse** (16 %, comme le temps clair) ; pluie/couvert ~22-23 %",
    "**Les fêtes ne s'emballent pas** : Noël = 15,5 % de graves vs 19,5 % en période normale",
    "-> météo & calendrier = **mesures dynamiques** (panneaux à messages variables), pas du béton : le budget rénovation reste sur le structurel",
], Inches(0.7), Inches(5.0), Inches(11.9), Inches(2.2), size=14, gap=8)

# S12 livrable
s = new_slide("Le livrable décisionnel - ce que le DOT ouvre lundi matin",
    notes="C'est ici que la valeur devient palpable : une ligne du CSV = une décision de travaux argumentée. La carte : couleur = aménagement recommandé, taille = risque réel par km.")
picture(s, img_path("00_carte"), Inches(0.55), Inches(1.55), width=Inches(6.4))
table(s, [
    ["Zone", "Accidents", "Aléa local", "Recommandation", "Impact estimé*"],
    ["#2", "2 543", "carrefour", "stop / cédez-le-passage", "≈ 505 graves évités"],
    ["#3", "2 431", "carrefour", "stop / cédez-le-passage", "≈ 483"],
    ["#8", "2 323", "carrefour", "stop / cédez-le-passage", "≈ 461"],
], Inches(7.25), Inches(1.7), Inches(5.6), col_widths=[0.9, 1.3, 1.4, 2.6, 2.2], font=11)
bullets(s, [
    "**zones_prioritaires.csv** : 50 zones triées par impact = ordre d'intervention",
    "**Carte interactive** : couleur = aménagement, taille = risque par km",
    "Impact cumulé top 50 : **≈ 6 860 accidents graves** potentiellement évités*",
    ("*ex-ante, associationnel - à valider par visite terrain puis pilote", 1),
], Inches(7.25), Inches(3.75), Inches(5.6), Inches(3.2), size=13.5, gap=9)

# S13 limites
s = new_slide("Limites & recul - ce qu'on ne prétend pas",
    notes="Annoncer les limites nous-mêmes = crédibilité. La grille valorise explicitement « deuils énoncés ». Diff-in-diff : zones traitées vs témoins comparables, avant/après travaux.")
bullets(s, [
    "**Pas de causalité** : aucune donnée avant/après installation -> signaux de priorisation, pas preuves d'efficacité",
    "**Biais de reporting** : couverture inégale - fiable sur les territoires à fort volume seulement",
    "Pas de profil conducteur (vitesse, alcool) , `severity` = seul proxy du coût",
    "**Exposition OSM partielle** (42/50) : API Overpass limitée - cache livré, dégradation propre",
], Inches(0.7), Inches(1.45), Inches(11.9), Inches(2.9), size=16, gap=11)
band(s, "**Mesurer l'impact réel (proposition)** : pilote sur 5-10 zones + **différence de différences** "
        "(zones traitées vs témoins, avant/après) - KPI : coût par accident grave évité.",
     Inches(0.7), Inches(4.7), Inches(11.9), height=Inches(1.15), accent=GREEN, size=15)

# S14 conclusion
s = new_slide("Conclusion - la valeur créée",
    notes="Finir sur « que fait le DOT lundi matin » et ouvrir sur le pilote terrain.")
tb = s.shapes.add_textbox(Inches(0.7), Inches(1.5), Inches(11.9), Inches(1.0))
add_rich(tb.text_frame.paragraphs[0],
         "On transforme un budget limité en **liste d'interventions priorisée, justifiée et explicable**.",
         22, bold_color=RED)
bullets(s, [
    "**OÙ** : 50 zones validées - concentration ×62, stabilité ρ = 0,78, exposition contrôlée",
    "**QUOI** : l'aménagement adapté à chaque zone + impact en accidents graves évités",
    "**QUAND** : mesures dynamiques vs travaux - bien séparés, budget préservé",
    "**Fiable** : Run All sans intervention, tests (notebook + pytest), versions figées, Docker",
], Inches(0.8), Inches(2.75), Inches(11.7), Inches(3.0), size=17, gap=13)
band(s, "Prochaine étape : un **pilote terrain** sur 5-10 zones pour transformer les associations en effets mesurés.",
     Inches(0.7), Inches(6.05), Inches(11.9), accent=GREEN, size=15)

# S15 séparateur annexes
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(1, 0, 0, SW, SH)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
tb = s.shapes.add_textbox(Inches(0.7), Inches(3.1), Inches(11.9), Inches(1.4))
p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
add_rich(p, "Annexes - Q&A", 40, color=WHITE, bold_color=GOLD)

# annexes
s = new_slide("Annexe A - Nettoyage & prétraitement (explicite, justifié)")
picture(s, img_path("01_a0"), Inches(0.55), Inches(1.55), width=Inches(6.6))
bullets(s, [
    "`end_lat/lng` (44 % NA) -> **abandonnées** ; `wind_chill` -> supprimée (redondante)",
    "Météo : imputation médiane , texte libre regroupé en 6 familles",
    "Booléens d'infra NA -> **absent** , sévérité binarisée (classe 2 = 80 %)",
    "**Tests d'intégrité bloquants** après nettoyage (6 assertions)",
], Inches(7.45), Inches(1.8), Inches(5.4), Inches(4.4), size=13.5, gap=10)

s = new_slide("Annexe B - Le signal descriptif brut (avant modèle)",
    notes="Ce delta brut est TROMPEUR (confondants) - c'est exactement ce que les effets ajustés corrigent. Bon exemple pédagogique si question.")
picture(s, img_path("04_a2bis"), Inches(0.55), Inches(1.55), width=Inches(7.0))
bullets(s, [
    "Sévérité moyenne **avec/sans** chaque élément (US entier)",
    "Lecture brute piégée : les équipements « marquent » les lieux denses",
    "-> c'est pourquoi la preuve finale utilise les **effets ajustés**",
], Inches(7.85), Inches(1.9), Inches(5.0), Inches(3.6), size=13.5, gap=10)

s = new_slide("Annexe C - Analyse des erreurs du modèle",
    notes="Au seuil 0,5 : 44 % des graves ratés. Abaisser le seuil ↑ rappel au prix de fausses alertes - arbitrage documenté par la courbe PR.")
picture(s, img_path("10_c2bis"), Inches(0.55), Inches(1.6), width=Inches(8.6))
bullets(s, [
    "Faux négatifs = **points noirs ratés** (44 % au seuil 0,5)",
    "Courbes PR vs plancher : l'arbitrage seuil <-> fausses alertes est explicite",
], Inches(9.35), Inches(2.0), Inches(3.4), Inches(3.6), size=13, gap=10)

s = new_slide("Annexe D - Robustesse : SHAP & convergence entre approches",
    notes="SHAP = requête d'attribution (corrélations du réel), g-computation = requête d'intervention. Les signes peuvent différer - documenté dans le notebook.")
picture(s, img_path("13_c4"), Inches(0.55), Inches(1.6), width=Inches(6.3))
picture(s, img_path("14_c4"), Inches(7.05), Inches(1.6), width=Inches(5.8))
bullets(s, [
    "Logistique, arbre, forêt et SHAP classent les leviers dans le **même ordre** (corrélations de rang élevées)",
    "Entrées : 12 booléens d'infra (leviers) + contrôles météo/temps/géo - **anti-fuite testé**",
], Inches(0.7), Inches(5.75), Inches(6.1), Inches(1.5), size=12.5, gap=7)

out = ROOT / "presentation_finale.pptx"
prs.save(str(out))
print(f"OK - {out.name} : {len(prs.slides.__iter__.__self__._sldIdLst)} slides")
